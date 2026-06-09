from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
import zipfile
from base64 import b64decode
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "scripts" / "build_report.py"
ONE_PIXEL_PNG = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


class BuildReportFormatTest(unittest.TestCase):
    def test_default_template_generates_formal_cover_numbered_headings_and_reference_page(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            draft = work / "report-draft.md"
            refs = work / "references.md"
            out = work / "report.docx"

            draft.write_text(
                "\n".join(
                    [
                        "# 绪论",
                        "这是第一章正文。",
                        "## 研究背景",
                        "这里是研究背景。",
                        "# 模型发展",
                        "这是第二章正文。",
                        "## 卷积网络",
                        "这里是卷积网络。",
                        "{{REFERENCES}}",
                    ]
                ),
                encoding="utf-8",
            )
            refs.write_text("[1] 测试参考文献。", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(BUILD),
                    "--draft",
                    str(draft),
                    "--refs",
                    str(refs),
                    "--output",
                    str(out),
                    "--title",
                    "深度学习模型发展历程与最新架构调研",
                    "--course",
                    "深度学习",
                    "--date",
                    "2026 年 6 月 9 日",
                ],
                cwd=ROOT,
                check=True,
            )

            doc = Document(str(out))
            text = "\n".join(p.text for p in doc.paragraphs)

            self.assertIn("报告题目    深度学习模型发展历程与最新架构调研", text)
            self.assertNotIn("放置", text)
            self.assertNotIn("校徽", text)
            self.assertNotIn("《XXXX》", text)
            self.assertNotIn("实验题目", text)

            self.assertIn("第一章 绪论", text)
            self.assertIn("1.1 研究背景", text)
            self.assertIn("第二章 模型发展", text)
            self.assertIn("2.1 卷积网络", text)

            with zipfile.ZipFile(out) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")

            self.assertIn("TOC \\o", xml)
            self.assertIn("1-3", xml)
            self.assertIn("\\h \\z \\u", xml)
            reference_pos = xml.index("参考文献")
            before_reference = xml[max(0, reference_pos - 1200) : reference_pos]
            self.assertRegex(before_reference, r'<w:br w:type="page"|<w:lastRenderedPageBreak')

    def test_figures_are_numbered_and_source_lines_hidden_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            assets = work / "assets"
            assets.mkdir()
            (assets / "figure.png").write_bytes(ONE_PIXEL_PNG)
            draft = work / "report-draft.md"
            refs = work / "references.md"
            out = work / "report.docx"

            draft.write_text(
                "\n".join(
                    [
                        "# 摘要",
                        "摘要正文。",
                        "# 引言",
                        "引言正文。",
                        "![概念图：结构演化](assets/figure.png)",
                        "图片来源：不应写入正文。",
                        "图片宽度：2cm",
                        "{{REFERENCES}}",
                    ]
                ),
                encoding="utf-8",
            )
            refs.write_text("[1] 测试参考文献。", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(BUILD),
                    "--draft",
                    str(draft),
                    "--refs",
                    str(refs),
                    "--output",
                    str(out),
                    "--root",
                    str(work),
                    "--title",
                    "图注测试",
                    "--course",
                    "深度学习",
                ],
                cwd=ROOT,
                check=True,
            )

            doc = Document(str(out))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("图2.1 概念图：结构演化", text)
            self.assertNotIn("图片来源：不应写入正文。", text)


if __name__ == "__main__":
    unittest.main()
