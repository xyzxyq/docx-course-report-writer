from base64 import b64decode
from io import BytesIO
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_report import replace_text_placeholders

PNG = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def add_field(paragraph, result="1"):
    for kind in ("begin", "separate", "end"):
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(node)
        if kind == "begin":
            instruction = OxmlElement("w:instrText")
            instruction.text = " PAGE "
            paragraph.add_run()._r.append(instruction)
        elif kind == "separate":
            paragraph.add_run(result)


class TemplatePlaceholderTest(unittest.TestCase):
    def test_split_marker_preserves_formatting_drawing_field_and_roundtrip(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("姓名：").bold = True
        p.add_run("{{STUDENT_").italic = True
        p.add_run("NAME}}").italic = True
        tail = p.add_run("（请核对）")
        tail.underline = True
        tail.add_picture(BytesIO(PNG))
        add_field(p)
        before_tail = etree.tostring(tail._r)
        field_xml = [etree.tostring(r._r) for r in p.runs[4:]]

        replace_text_placeholders(doc, {"STUDENT_NAME": "张三"})

        self.assertEqual("姓名：张三（请核对）1", p.text)
        self.assertEqual(before_tail, etree.tostring(tail._r))
        self.assertEqual(field_xml, [etree.tostring(r._r) for r in p.runs[4:]])
        self.assertTrue(p.runs[1].italic)
        self.assertEqual("张三", p.runs[1].text)
        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        reopened = Document(stream)
        self.assertEqual(1, len(reopened.inline_shapes))
        with zipfile.ZipFile(stream) as package:
            media = [name for name in package.namelist() if name.startswith("word/media/")]
            self.assertEqual([PNG], [package.read(name) for name in media])
        self.assertEqual(p.text, reopened.paragraphs[0].text)
        original_xml = etree.tostring(reopened._element)
        replace_text_placeholders(reopened, {"STUDENT_NAME": "张三"})
        self.assertEqual(original_xml, etree.tostring(reopened._element))

    def test_repeated_markers_preserve_prefix_suffix_and_same_run_image(self):
        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("前 {{NAME}} / [[NAME]] / 【NAME】 / {{NAME}} 后")
        run.bold = True
        run.add_picture(BytesIO(PNG))
        drawing = etree.tostring(run._r.xpath("w:drawing")[0])
        replace_text_placeholders(doc, {"NAME": " A & <B> "})
        self.assertEqual("前  A & <B>  /  A & <B>  /  A & <B>  /  A & <B>  后", p.text)
        self.assertTrue(run.bold)
        self.assertEqual(drawing, etree.tostring(run._r.xpath("w:drawing")[0]))
        standalone = doc.add_paragraph("{{NAME}}")
        replace_text_placeholders(doc, {"NAME": " A & <B> "})
        self.assertEqual(" A & <B> ", standalone.text)
        self.assertEqual("preserve", standalone._p.xpath(".//w:t")[0].get(qn("xml:space")))

    def test_table_label_respects_explicit_marker_and_plain_cell_image(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for row in table.rows:
            row.cells[0].text = "学生姓名"
        p = table.cell(0, 1).paragraphs[0]
        p.add_run("{{STUDENT_").italic = True
        p.add_run("NAME}}")
        p.add_run("（组长）").underline = True
        p.add_run().add_picture(BytesIO(PNG))
        plain = table.cell(1, 1).paragraphs[0]
        plain.add_run("旧姓名").italic = True
        plain.add_run().add_picture(BytesIO(PNG))
        replace_text_placeholders(doc, {"学生姓名": "张三", "STUDENT_NAME": "张三"})
        self.assertEqual("张三（组长）", p.text)
        self.assertTrue(p.runs[2].underline)
        self.assertEqual("张三", plain.text)
        self.assertTrue(plain.runs[0].italic)
        self.assertEqual(2, len(doc.inline_shapes))

    def test_label_append_and_empty_cell_preserve_existing_objects(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("学生姓名").bold = True
        p.add_run().add_picture(BytesIO(PNG))
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "学生姓名"
        table.cell(0, 1).paragraphs[0].add_run().add_picture(BytesIO(PNG))
        replace_text_placeholders(doc, {"学生姓名": "张三"})
        self.assertEqual("学生姓名    张三", p.text)
        self.assertEqual("张三", table.cell(0, 1).text)
        self.assertEqual(2, len(doc.inline_shapes))

    def test_multiline_replacement_keeps_tabs_breaks_and_surrounding_text(self):
        doc = Document()
        p = doc.add_paragraph("说明：{{NOTE}}。")
        replace_text_placeholders(doc, {"NOTE": "第一行\n第二行\t值"})
        self.assertEqual("说明：第一行\n第二行\t值。", p.text)
        self.assertEqual(1, len(p._p.xpath(".//w:br")))
        self.assertEqual(1, len(p._p.xpath(".//w:tab")))

    def test_marker_inside_field_or_across_complex_boundary_is_rejected(self):
        for kind in ("field", "drawing", "hyperlink", "revision"):
            with self.subTest(kind=kind):
                doc = Document()
                p = doc.add_paragraph()
                if kind == "field":
                    add_field(p, "{{NAME}}")
                else:
                    p.add_run("{{NA")
                    if kind == "drawing":
                        p.add_run().add_picture(BytesIO(PNG))
                    else:
                        p._p.append(OxmlElement("w:hyperlink" if kind == "hyperlink" else "w:ins"))
                    p.add_run("ME}}")
                before = etree.tostring(p._p)
                with self.assertRaisesRegex(ValueError, "complex|boundary|field"):
                    replace_text_placeholders(doc, {"NAME": "张三"})
                self.assertEqual(before, etree.tostring(p._p))

    def test_merged_label_cell_is_not_treated_as_its_own_value_cell(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        cell.text = "学生姓名"
        replace_text_placeholders(doc, {"学生姓名": "张三"})
        self.assertEqual("学生姓名", cell.text)

    def test_field_spanning_paragraphs_is_not_edited(self):
        doc = Document()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        doc.add_paragraph().add_run()._r.append(begin)
        doc.add_paragraph("{{NAME}}")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        doc.add_paragraph().add_run()._r.append(end)
        before = etree.tostring(doc._element)
        with self.assertRaisesRegex(ValueError, "complex"):
            replace_text_placeholders(doc, {"NAME": "张三"})
        self.assertEqual(before, etree.tostring(doc._element))

    def test_word_proofing_markers_do_not_break_plain_text_replacement(self):
        doc = Document()
        p = doc.add_paragraph("{{NA")
        start = OxmlElement("w:proofErr")
        start.set(qn("w:type"), "spellStart")
        p._p.append(start)
        p.add_run("ME")
        end = OxmlElement("w:proofErr")
        end.set(qn("w:type"), "spellEnd")
        p._p.append(end)
        p.add_run("}}")
        before = [etree.tostring(node) for node in p._p.xpath("w:proofErr")]
        replace_text_placeholders(doc, {"NAME": "张三"})
        self.assertEqual("张三", p.text)
        self.assertEqual(before, [etree.tostring(node) for node in p._p.xpath("w:proofErr")])

    def test_every_run_split_replaces_multiple_markers_without_touching_link(self):
        source = "前{{NAME}}中{{NAME}}后"
        for split in range(len(source) + 1):
            with self.subTest(split=split):
                doc = Document()
                p = doc.add_paragraph()
                p.add_run(source[:split]).bold = True
                p.add_run(source[split:]).italic = True
                link = OxmlElement("w:hyperlink")
                link.set(qn("w:anchor"), "bookmark")
                run = OxmlElement("w:r")
                run.text = "附注链接"
                link.append(run)
                p._p.append(link)
                original = etree.tostring(link)
                replace_text_placeholders(doc, {"NAME": "张三"})
                self.assertEqual("前张三中张三后附注链接", p.text)
                self.assertEqual(original, etree.tostring(link))

    def test_builder_preserves_template_assets_in_saved_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("姓名：").bold = True
            p.add_run("{{STUDENT_NAME}}")
            p.add_run("（请核对）").underline = True
            p.add_run().add_picture(BytesIO(PNG))
            add_field(p)
            template = root / "template.docx"
            doc.save(template)
            original = template.read_bytes()
            draft = root / "draft.md"
            draft.write_text("# 测试\n正文。", encoding="utf-8")
            refs = root / "refs.md"
            refs.write_text("", encoding="utf-8")
            out = root / "report.docx"
            subprocess.run(
                [sys.executable, str(ROOT / "scripts/build_report.py"),
                 "--draft", str(draft), "--refs", str(refs),
                 "--template", str(template), "--student-name", "张三",
                 "--output", str(out), "--keep-template-body", "--no-toc"],
                check=True, capture_output=True, text=True,
            )
            result = Document(out)
            self.assertEqual("姓名：张三（请核对）1", result.paragraphs[0].text)
            self.assertTrue(result.paragraphs[0].runs[2].underline)
            self.assertEqual(1, len(result.inline_shapes))
            self.assertEqual(3, len(result.paragraphs[0]._p.xpath(".//w:fldChar")))
            self.assertEqual(original, template.read_bytes())

    def test_builder_rejects_unsafe_template_without_overwriting_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            doc = Document()
            add_field(doc.add_paragraph(), "{{STUDENT_NAME}}")
            doc.save(template)
            original = template.read_bytes()
            draft = root / "draft.md"
            draft.write_text("# 测试\n正文。", encoding="utf-8")
            refs = root / "refs.md"
            refs.write_text("", encoding="utf-8")
            out = root / "report.docx"
            out.write_bytes(b"previous checked report")
            result = subprocess.run(
                [sys.executable, str(ROOT / "scripts/build_report.py"),
                 "--draft", str(draft), "--refs", str(refs), "--template", str(template),
                 "--student-name", "张三", "--output", str(out), "--no-toc"],
                capture_output=True, text=True,
            )
            self.assertNotEqual(0, result.returncode)
            self.assertIn("complex", result.stderr)
            self.assertEqual(b"previous checked report", out.read_bytes())
            self.assertEqual(original, template.read_bytes())


if __name__ == "__main__":
    unittest.main()
