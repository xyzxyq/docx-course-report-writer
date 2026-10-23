from io import BytesIO
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.shared import Cm
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from vector_figure import ASVG_NS, add_vector_picture, validate_svg, vector_assets

try:
    import pymupdf
except ImportError:
    pymupdf = None

SVG = b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"><rect x="5" y="5" width="190" height="90" fill="#abcdef"/></svg>'
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships", "asvg": ASVG_NS}


class SvgBoundaryTest(unittest.TestCase):
    def test_external_and_active_content_rejected(self):
        for content in (
            '<image href="file:///private/test.png"/>',
            '<use href="https://example.com/a.svg#x"/>',
            '<script>doSomething()</script>',
            '<rect onload="doSomething()"/>',
            '<style>@import "https://example.com/a.css";</style>',
            '<rect fill="url(https://example.com/a.svg)"/>',
            '<foreignObject/>',
        ):
            with self.subTest(content=content), self.assertRaises(ValueError):
                validate_svg(f'<svg xmlns="http://www.w3.org/2000/svg">{content}</svg>'.encode())
        with self.assertRaises(ValueError):
            validate_svg(b'<!DOCTYPE svg [<!ENTITY x "expanded">]>' + SVG)

    def test_local_glyph_reference_is_supported(self):
        validate_svg(b'<svg xmlns="http://www.w3.org/2000/svg"><defs><path id="g" d="M0 0L5 5"/></defs><use href="#g"/></svg>')
        validate_svg(b'<svg xmlns="http://www.w3.org/2000/svg"><text>C:\\data</text></svg>')

    def test_missing_dependency_has_actionable_error(self):
        with patch.dict(sys.modules, {"pymupdf": None}), self.assertRaisesRegex(RuntimeError, "pip install PyMuPDF"):
            vector_assets(Path("unused.pdf"))


@unittest.skipUnless(pymupdf, "PyMuPDF required for vector integration tests")
class VectorFigureTest(unittest.TestCase):
    def test_svg_parts_and_multiple_relationships_survive_docx_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fig = root / "图.svg"
            fig.write_bytes(SVG)
            doc = Document()
            for _ in range(2):
                pic = add_vector_picture(doc.add_paragraph().add_run(), fig, Cm(10))
                self.assertAlmostEqual(pic.width / pic.height, 2, places=2)
            out = root / "report.docx"
            doc.save(out)
            Document(out).save(out)
            with zipfile.ZipFile(out) as z:
                xml = etree.fromstring(z.read("word/document.xml"))
                rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
                targets = {r.get("Id"): r.get("Target") for r in rels}
                for blip in xml.findall(".//a:blip", NS):
                    fallback = targets[blip.get(f'{{{NS["r"]}}}embed')]
                    self.assertTrue(z.read("word/" + fallback).startswith(b"\x89PNG"))
                    svg_blip = blip.find(".//asvg:svgBlip", NS)
                    target = targets[svg_blip.get(f'{{{NS["r"]}}}embed')]
                    self.assertEqual(z.read("word/" + target), SVG)
                self.assertEqual(len(xml.findall(".//asvg:svgBlip", NS)), 2)
                self.assertIn(b"image/svg+xml", z.read("[Content_Types].xml"))

    def test_real_pdf_converts_to_paths_and_is_inserted_by_builder(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf = pymupdf.open()
            page = pdf.new_page(width=200, height=100)
            page.draw_rect((5, 5, 195, 95), color=(0, 0, 1))
            page.insert_text((20, 50), "Vector fixture")
            path = root / "figure.PDF"
            pdf.save(path)
            pdf.close()
            svg, png = vector_assets(path)
            self.assertIn(b"<path", svg)
            self.assertNotIn(b"<image", svg)
            self.assertTrue(png.startswith(b"\x89PNG"))
            (root / "draft.md").write_text("# 测试\n![矢量图](figure.PDF)\n", encoding="utf-8")
            (root / "refs.md").write_text("", encoding="utf-8")
            subprocess.run([sys.executable, str(ROOT / "scripts/build_report.py"), "--draft", str(root / "draft.md"), "--refs", str(root / "refs.md"), "--root", str(root), "--output", str(root / "report.docx"), "--title", "矢量插图测试"], check=True, capture_output=True)
            with zipfile.ZipFile(root / "report.docx") as z:
                self.assertIn(b"svgBlip", z.read("word/document.xml"))
                self.assertIn("图1.1 矢量图", Document(BytesIO((root / "report.docx").read_bytes())).paragraphs[-1].text)

    def test_multipage_pdf_rejected_without_changing_previous_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf = pymupdf.open()
            pdf.new_page(); pdf.new_page()
            pdf.save(root / "two.pdf")
            pdf.close()
            with self.assertRaisesRegex(ValueError, "exactly one page"):
                vector_assets(root / "two.pdf")
            (root / "draft.md").write_text("# Test\n![figure](two.pdf)\n")
            (root / "refs.md").write_text("")
            out = root / "report.docx"
            out.write_bytes(b"previous approved report")
            result = subprocess.run([sys.executable, str(ROOT / "scripts/build_report.py"), "--draft", str(root / "draft.md"), "--refs", str(root / "refs.md"), "--root", str(root), "--output", str(out), "--title", "Test"], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(out.read_bytes(), b"previous approved report")


if __name__ == "__main__":
    unittest.main()
